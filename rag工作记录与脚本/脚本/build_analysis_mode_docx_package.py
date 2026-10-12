from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE = Path(r"D:\studyprogram\mypro\rag工作记录与脚本\评测产物与工作目录\analysis-mode-only-tender-2026-09-19")
OUT = BASE / "analysis-mode-only-tender-2026-09-19.docx"

DOCS = [
    ("总入口说明", BASE / "README.md"),
    ("投标人可见", BASE / "01_投标人可见" / "TenderAnalysisSummary_投标分析摘要.md"),
    ("投标人可见", BASE / "01_投标人可见" / "BidRiskChecklist_投标风险清单.md"),
    ("投标人可见", BASE / "01_投标人可见" / "ScoringOpportunityReport_得分机会报告.md"),
    ("投标人可见", BASE / "01_投标人可见" / "MissingMaterialsList_缺失材料清单.md"),
    ("投标人可见", BASE / "01_投标人可见" / "CollaborationAgenda_协作议程.md"),
    ("投标人可见", BASE / "01_投标人可见" / "ResponsePlanPreview_生成前响应计划预览.md"),
    ("人工决策", BASE / "02_人工决策" / "ManualDecisionBoard_人工决策板.md"),
    ("人工决策", BASE / "02_人工决策" / "ConfirmationRequest_确认请求.md"),
    ("下一阶段中间产物", BASE / "03_下一阶段中间产物" / "TenderBrief.md"),
    ("下一阶段中间产物", BASE / "03_下一阶段中间产物" / "RequirementMatrix.md"),
    ("下一阶段中间产物", BASE / "03_下一阶段中间产物" / "ScoringMatrix.md"),
    ("下一阶段中间产物", BASE / "03_下一阶段中间产物" / "SubmissionStructurePack.md"),
    ("下一阶段中间产物", BASE / "03_下一阶段中间产物" / "PricingRulePack.md"),
    ("下一阶段中间产物", BASE / "03_下一阶段中间产物" / "FatalRiskRulePack.md"),
    ("下一阶段中间产物", BASE / "03_下一阶段中间产物" / "EvidenceAvailabilityPack.md"),
    ("下一阶段中间产物", BASE / "03_下一阶段中间产物" / "ScoreProjectionPack.md"),
    ("下一阶段中间产物", BASE / "03_下一阶段中间产物" / "GapList.md"),
    ("下一阶段中间产物", BASE / "03_下一阶段中间产物" / "StrategyCandidatePack.md"),
    ("下一阶段中间产物", BASE / "03_下一阶段中间产物" / "StrategyEvaluationPack.md"),
    ("下一阶段中间产物", BASE / "03_下一阶段中间产物" / "AnalysisReviewReport.md"),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, margin_twips: int = 100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_twips))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, pct: int = 5000) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(pct))
    tbl_w.set(qn("w:type"), "pct")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)

    for style_name, size in (("Title", 18), ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)

    return doc


def add_text_with_inline_code(paragraph, text: str, size: float = 10.5, bold: bool = False) -> None:
    pattern = re.compile(r"`([^`]+)`")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            set_run_font(run, size=size, bold=bold)
        run = paragraph.add_run(match.group(1))
        set_run_font(run, size=size, bold=bold)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=size, bold=bold)


def split_table_row(line: str) -> list[str]:
    stripped = line.strip().strip("|")
    return [cell.strip() for cell in stripped.split("|")]


def is_separator_row(line: str) -> bool:
    cells = split_table_row(line)
    if not cells:
        return False
    return all(re.fullmatch(r":?-{3,}:?", c.strip()) for c in cells)


def collect_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
        if not is_separator_row(lines[i]):
            rows.append(split_table_row(lines[i]))
        i += 1
    return rows, i


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_width(table)
    for r_idx, row in enumerate(rows):
        if r_idx == 0:
            set_repeat_table_header(table.rows[r_idx])
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell, 95)
            if r_idx == 0:
                set_cell_shading(cell, "1F4E79")
            elif r_idx % 2 == 0:
                set_cell_shading(cell, "F3F6FA")
            text = row[c_idx] if c_idx < len(row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.12
            add_text_with_inline_code(paragraph, text, size=9.0, bold=(r_idx == 0))
            for run in paragraph.runs:
                if r_idx == 0:
                    run.font.color.rgb = RGBColor(255, 255, 255)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def add_markdown_content(doc: Document, path: Path, section_label: str, first_file: bool = False) -> None:
    text = path.read_text(encoding="utf-8")
    lines = text.splitlines()
    if not first_file:
        doc.add_page_break()
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(3)
    run = meta.add_run(section_label)
    set_run_font(run, size=9.5, bold=True)
    run.font.color.rgb = RGBColor(90, 90, 90)

    i = 0
    in_code = False
    code_buffer: list[str] = []
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        if line.strip().startswith("```"):
            if in_code:
                p = doc.add_paragraph(style=None)
                p.paragraph_format.left_indent = Cm(0.3)
                p.paragraph_format.space_after = Pt(6)
                run = p.add_run("\n".join(code_buffer))
                set_run_font(run, size=9.0)
                run.font.name = "Consolas"
                run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
                run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
                code_buffer = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buffer.append(line)
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        if line.strip().startswith("|") and line.strip().endswith("|"):
            rows, nxt = collect_table(lines, i)
            add_table(doc, rows)
            i = nxt
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            content = heading.group(2).strip()
            if level == 1:
                p = doc.add_paragraph(style="Heading 1")
            elif level == 2:
                p = doc.add_paragraph(style="Heading 2")
            else:
                p = doc.add_paragraph(style="Heading 3")
            add_text_with_inline_code(p, content, size={1: 15, 2: 13, 3: 11.5, 4: 11.5}.get(level, 11.5), bold=True)
            i += 1
            continue
        bullet = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_text_with_inline_code(p, bullet.group(2).strip(), size=10.2)
            i += 1
            continue
        numbered = re.match(r"^(\s*)\d+\.\s+(.*)$", line)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            add_text_with_inline_code(p, numbered.group(2).strip(), size=10.2)
            i += 1
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        add_text_with_inline_code(p, line.strip(), size=10.5)
        i += 1


def add_cover(doc: Document) -> None:
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text_with_inline_code(title, "广州供电局车辆维保项目分析模式结论包", size=18, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text_with_inline_code(subtitle, "仅招标文件输入  生成前分析与协作确认产物", size=11)
    subtitle.paragraph_format.space_after = Pt(18)

    intro = doc.add_paragraph()
    intro.paragraph_format.line_spacing = 1.2
    add_text_with_inline_code(
        intro,
        "本 Word 文档汇总分析模式的全部可交付产物。输入边界限定为真实招标文件，不使用中标样本、历史投标文件或投标人企业资料。因此本文给出的是规则、风险、证据缺口、人工确认点和下一阶段门禁，不给出某家投标人的确定得分。",
        size=10.8,
    )

    overview = [
        ["分区", "内容", "读者"],
        ["投标人可见", "投标分析摘要、风险清单、得分机会、缺失材料、协作议程、生成前响应计划", "投标人、业务负责人、标书编制人员"],
        ["人工决策", "人工决策板、确认请求", "项目负责人、报价负责人、商务和技术负责人"],
        ["下一阶段中间产物", "TenderBrief、矩阵、规则包、缺口、策略评估、分析复核报告", "协作模式、生成模式、研发质检"],
    ]
    add_table(doc, overview)


def add_update_fields_setting(doc: Document) -> None:
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def add_field_run(paragraph, field_type: str, text: str | None = None) -> None:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), field_type)
    run._r.append(fld_char)
    if text:
        result = paragraph.add_run(text)
        set_run_font(result, size=10.5)


def add_toc_instruction(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run._r.append(instr)


def add_automatic_toc(doc: Document) -> None:
    doc.add_page_break()
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    heading.paragraph_format.space_after = Pt(10)
    add_text_with_inline_code(heading, "目录", size=15, bold=True)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(8)
    add_text_with_inline_code(
        intro,
        "本页使用 Word 自动目录字段生成。打开文档后如目录未显示页码，请在 Word 中更新目录字段。",
        size=10.5,
    )

    toc = doc.add_paragraph()
    toc.paragraph_format.space_after = Pt(0)
    add_field_run(toc, "begin")
    add_toc_instruction(toc, r'TOC \o "1-3" \h \z \u')
    add_field_run(toc, "separate", "请在 Word 中右键更新域以生成目录。")
    add_field_run(toc, "end")


def add_footer(doc: Document) -> None:
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("分析模式结论包  仅招标文件输入")
        set_run_font(run, size=8.5)
        run.font.color.rgb = RGBColor(100, 100, 100)


def main() -> None:
    doc = setup_document()
    add_update_fields_setting(doc)
    add_cover(doc)
    add_automatic_toc(doc)
    for idx, (section_label, path) in enumerate(DOCS):
        add_markdown_content(doc, path, section_label, first_file=False)
    add_footer(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
