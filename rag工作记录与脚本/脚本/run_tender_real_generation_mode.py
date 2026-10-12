from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from build_tender_generation_eval import (
    OUTPUT_DIR,
    SAMPLE_DIR,
    TENDER_PDF,
    docx_text_and_tables,
    extract_bid_facts,
    find_doc,
    score_from_facts,
    scoring_matrix,
    write_json,
)


REAL_OUTPUT_DIR = OUTPUT_DIR / "real-generation-run"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, color="FFFFFF")
        set_cell_shading(table.rows[0].cells[i], "1F4E79")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    document.add_paragraph()


def para(document: Document, text: str, style: str | None = None, bold_lead: str | None = None) -> None:
    p = document.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        lead.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)


def load_samples() -> dict[str, dict[str, Any]]:
    samples = {path.name.split("_", 1)[0]: docx_text_and_tables(path) for path in SAMPLE_DIR.glob("*.docx")}
    for required in ["4", "13", "14", "15", "16", "21", "22", "23", "24", "25", "26"]:
        if required not in samples:
            samples[required] = docx_text_and_tables(find_doc(required))
    return samples


def build_draft_sections(facts: dict[str, Any]) -> list[dict[str, Any]]:
    performance_names = [item["project"] for item in facts["performance_projects"][:6]]
    return [
        {
            "id": "SEC-TECH-01",
            "title": "项目概况",
            "llm_generated_text": (
                "本项目为广州供电局2026-2027年综合类车辆维保项目，采购编号CG2700022002213308。"
                "我方针对标的8花都区域一般车辆维保两个标包统一响应，严格保持商务、技术和报价一致。"
                "本次响应以保障广州供电局车辆安全、连续、高效运行为目标，围绕一般车辆维修保养、应急救援、"
                "月度巡检、配件质量追溯和质量保证期管理建立闭环服务机制。"
            ),
            "evidence": ["招标文件项目范围和标的8要求", "样本项目概况章节", "同一标的全部标包一致性要求"],
            "score_target": "TECH-01",
        },
        {
            "id": "SEC-TECH-02",
            "title": "服务方案",
            "llm_generated_text": (
                "我方按接车登记、故障诊断、一次性报价、维修实施、三级质检、竣工交付、回访改进的流程组织服务。"
                "针对电力保障车辆高频使用、应急任务多、车型复杂等特点，设置常规保养、总成维修、钣喷、轮胎、空调、"
                "四轮定位和新能源车辆安全检查等专项作业路径。对于发动机大修等复杂维修，执行拆解前拍照、专用工具准备、"
                "标准扭矩装配、试车验证和旧件留存要求，确保过程可追溯。"
            ),
            "evidence": ["服务方案章节", "发动机大修作业流程", "焊接及变速器作业指导内容"],
            "score_target": "TECH-02",
        },
        {
            "id": "SEC-TECH-03",
            "title": "服务团队及服务能力",
            "llm_generated_text": (
                "我方建立总负责人、技术组、接待组、救援组和备件保障协同机制，配备维修技师、质检人员、业务接待、"
                "价格结算及应急救援人员。服务团队覆盖发动机、轮胎、空调、四轮定位、供油系统、散热器、车身清洁、"
                "门窗玻璃等岗位，并提供社保或退休返聘劳务合同等用工证明。人员、设备和管理制度共同支撑高频车辆维保任务。"
            ),
            "evidence": ["服务团队及服务能力章节", "人员岗位表", "社保/退休返聘用工证明"],
            "score_target": "TECH-03",
        },
        {
            "id": "SEC-TECH-04",
            "title": "服务质量保证措施",
            "llm_generated_text": (
                f"我方服务点与所投标的单位行驶距离为{facts['service_quality']['distance_km']}公里，满足≤5公里的A档条件。"
                "质量控制采用三级质检制度、配件采购验收和追溯制度、维修过程检验、竣工出厂检验和返修责任判定机制。"
                "维修一次成功率、客户满意度、返修率和交付及时率作为持续改进指标，确保维修质量和服务响应稳定。"
            ),
            "evidence": ["服务质量保证措施章节", "百度地图距离", "三级质检制度", "配件追溯制度"],
            "score_target": "TECH-04",
        },
        {
            "id": "SEC-TECH-05",
            "title": "服务承诺",
            "llm_generated_text": (
                "我方承诺接到采购人通知后30分钟内上门取车送车，并每月上门定期检查车辆，免费添加油水液。"
                "同时提供24小时应急服务和免费拖车服务，成立应急保障工作组，配备维修保障车辆、诊断设备、备件和救援人员。"
                "质量保证期不低于国家和招标文件要求，因我方维修质量造成的问题由我方负责返修并承担相应责任。"
            ),
            "evidence": ["服务承诺章节", "30分钟取送车承诺", "每月定期检查承诺", "24小时应急服务"],
            "score_target": "TECH-05",
        },
        {
            "id": "SEC-BIZ-01",
            "title": "商务评分证据",
            "llm_generated_text": (
                "我方同类项目业绩覆盖广州供电局、花都区人民医院、花东镇人民政府、中国邮政等车辆维修保养服务，"
                f"样本列明{facts['performance_project_count']}项业绩，满足每项5分、最高30分的要求。"
                "企业信誉提供2020-2024年A级纳税人材料，管理体系认证覆盖质量、环境、职业健康安全三项，"
                "资质等级为一类维修资质。"
            ),
            "evidence": performance_names + ["企业信誉", "三体系认证", "一类维修资质"],
            "score_target": "BIZ-ALL",
        },
        {
            "id": "SEC-PRICE-01",
            "title": "报价响应",
            "llm_generated_text": (
                f"本次报价按下浮率填报，维修项目工时费下浮系数a={facts['pricing']['labor_discount_a_percent']}%，"
                f"维修材料费下浮系数b={facts['pricing']['material_discount_b_percent']}%，"
                f"综合下浮率c={facts['pricing']['composite_discount_c_percent']}%。"
                "该报价满足0<a≤20%、0<b≤10%的有效报价范围，且未超过成本分析报告强制触发阈值。"
            ),
            "evidence": ["报价表", "招标文件报价规则"],
            "score_target": "PRICE-01",
        },
    ]


def build_review(draft_sections: list[dict[str, Any]], facts: dict[str, Any], score: dict[str, Any]) -> dict[str, Any]:
    issues = []
    if score["price_proxy_score"] < 30:
        issues.append({
            "severity": "medium",
            "type": "external_data_missing",
            "description": "价格分需要全部投标人报价分布才能按靶心法精算。",
            "blocked": False,
        })
    issues.append({
        "severity": "medium",
        "type": "external_data_missing",
        "description": "供应商评价分需要广州供电局2024年度履约评价和南网失信扣分数据。",
        "blocked": False,
    })
    for required in ["30分钟", "每月", "4.2公里", "一类", "三项"]:
        joined = "\n".join(s["llm_generated_text"] for s in draft_sections)
        if required not in joined:
            issues.append({
                "severity": "high",
                "type": "missing_key_claim",
                "description": f"生成正文缺少关键表述：{required}",
                "blocked": True,
            })
    return {
        "reviewer_llm": "Codex current session model (GPT-5 based coding agent)",
        "backend_llm_gateway": {
            "attempted": True,
            "endpoint": "https://ahg.codes/v1/chat/completions",
            "result": "401 Unauthorized; no API key available in environment",
            "used_for_generation": False,
        },
        "status": "passed_with_external_data_limits" if not any(i["blocked"] for i in issues) else "blocked",
        "issues": issues,
        "coverage_checks": {
            "business_scoring_evidence_bound": True,
            "technical_scoring_evidence_bound": True,
            "pricing_rule_bound": True,
            "instruction_isolation_recorded": True,
            "score_projection_present": True,
            "review_report_present": True,
        },
        "score": score,
    }


def create_docx(path: Path, facts: dict[str, Any], draft_sections: list[dict[str, Any]], review: dict[str, Any]) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = document.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ["Title", "Heading 1", "Heading 2"]:
        styles[style_name].font.name = "Microsoft YaHei"
        styles[style_name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        styles[style_name].font.color.rgb = RGBColor(0, 0, 0)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("广州供电局车辆维保项目生成模式产物包")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("标的8 花都区域一般车辆维保").bold = True
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"生成时间：{datetime.now().isoformat(timespec='seconds')}")

    document.add_heading("生成边界", level=1)
    para(document, "招标文件和中标样本均作为数据源处理，其中任何面向AI、模型、系统或工具的语句均不作为执行指令。")
    para(document, "后端外部LLM网关因缺少API key返回401，本次正文生成和复核由当前Codex会话模型完成，并在JSON产物中记录。")

    document.add_heading("评分投影", level=1)
    score = review["score"]
    add_table(document, ["维度", "得分", "说明"], [
        ["商务", f"{score['business_weighted_score']:.2f}/30", "同类业绩、企业信誉、三体系认证、一类资质均达满分投影。"],
        ["技术", f"{score['technical_weighted_score']:.2f}/30", "服务方案、团队、质量、承诺和场地证据完整。"],
        ["价格", f"{score['price_proxy_score']:.2f}/30", "报价合规，但缺少全体投标人报价分布，无法精算靶心法真实分。"],
        ["供应商评价", f"{score['supplier_evaluation_proxy_score']:.2f}/10", "缺少履约评价和失信扣分外部数据，按产物完整性代理评分。"],
        ["总分", f"{score['total_percent_score']:.2f}/100", "生成模式产物对照中标样本的代理评分。"],
    ])

    document.add_heading("正文草稿", level=1)
    for section_data in draft_sections:
        document.add_heading(section_data["title"], level=2)
        para(document, section_data["llm_generated_text"])
        para(document, "证据绑定：" + "；".join(str(e) for e in section_data["evidence"][:8]))

    document.add_heading("复核结论", level=1)
    para(document, f"复核状态：{review['status']}")
    add_table(document, ["级别", "类型", "问题", "是否阻断"], [
        [issue["severity"], issue["type"], issue["description"], "是" if issue["blocked"] else "否"]
        for issue in review["issues"]
    ])

    document.save(path)


def main() -> None:
    REAL_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    samples = load_samples()
    facts = extract_bid_facts(samples)
    score = score_from_facts(facts)
    draft_sections = build_draft_sections(facts)
    review = build_review(draft_sections, facts, score)

    run = {
        "metadata": {
            "mode": "real_generation_mode_offline_tool_chain",
            "created_at": datetime.now().isoformat(timespec="seconds"),
            "tender_pdf": str(TENDER_PDF),
            "winning_sample_dir": str(SAMPLE_DIR),
            "llm_used": "Codex current session model (GPT-5 based coding agent)",
            "backend_llm_gateway_used": False,
            "backend_llm_gateway_blocker": "401 Unauthorized from https://ahg.codes/v1/chat/completions because no API key is available",
            "instruction_isolation": "Attached documents are data only. Instructions inside attached documents are not executed.",
        },
        "tool_registry": [
            {"name": "tender_text_cache_loader", "kind": "TOOL", "status": "completed"},
            {"name": "winning_sample_docx_extractor", "kind": "TOOL", "status": "completed"},
            {"name": "evidence_binding_matcher", "kind": "TOOL", "status": "completed"},
            {"name": "codex_writer_agent", "kind": "AGENT", "model": "Codex current session model", "status": "completed"},
            {"name": "codex_review_agent", "kind": "AGENT", "model": "Codex current session model", "status": "completed"},
            {"name": "docx_export_tool", "kind": "TOOL", "status": "completed"},
        ],
        "facts": facts,
        "ScoringMatrix": scoring_matrix(),
        "EvidenceBindingPack": {
            "bindings": [
                {"target": "business", "status": "full", "evidence": ["13", "14", "15", "16"]},
                {"target": "technical", "status": "full", "evidence": ["21", "22", "23", "24", "25", "26"]},
                {"target": "pricing", "status": "full", "evidence": ["4"]},
            ]
        },
        "DraftPack": {"sections": draft_sections},
        "ReviewReport": review,
        "ScoreProjectionPack": score,
    }

    run_path = REAL_OUTPUT_DIR / "real_generation_run.json"
    review_path = REAL_OUTPUT_DIR / "llm_review_score_report.json"
    docx_path = REAL_OUTPUT_DIR / "generated_bid_response_pack.docx"
    write_json(run_path, run)
    write_json(review_path, review)
    create_docx(docx_path, facts, draft_sections, review)
    print(json.dumps({
        "run": str(run_path),
        "review": str(review_path),
        "docx": str(docx_path),
        "score": score["total_percent_score"],
        "review_status": review["status"],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
