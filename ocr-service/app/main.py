import io, json, logging, os, tempfile, uuid
from pathlib import Path
from fastapi import FastAPI, File, UploadFile, HTTPException, Query
from fastapi.responses import JSONResponse
from PIL import Image
import numpy as np

logging.basicConfig(level=logging.INFO)
log = logging.getLogger("agicore-ocr")
app = FastAPI(title="AGI Core Local OCR")
MAX_BYTES = 100 * 1024 * 1024
_ocr = _structure = None


def normalize_image(img):
    array = np.asarray(img)
    if array.ndim == 2:
        array = np.repeat(array[:, :, None], 3, axis=2)
    elif array.ndim == 3 and array.shape[2] == 1:
        array = np.repeat(array, 3, axis=2)
    elif array.ndim == 3 and array.shape[2] >= 3:
        array = array[:, :, :3]
    else:
        raise ValueError("image must be a 2D grayscale or 3D color array")
    if array.dtype != np.uint8:
        array = np.clip(array, 0, 255).astype(np.uint8)
    return np.ascontiguousarray(array)

def engines():
    global _ocr, _structure
    if _ocr is None:
        from paddleocr import PaddleOCR, PPStructureV3
        _ocr = PaddleOCR(use_doc_orientation_classify=False, use_doc_unwarping=False,
                         use_textline_orientation=False, engine="paddle", enable_mkldnn=False)
        # CPU profile: keep layout/OCR/table structure only. Optional modules are
        # explicitly disabled because their default models are unnecessary for RAG tables.
        _structure = PPStructureV3(
            use_doc_orientation_classify=False,
            use_doc_unwarping=False,
            use_textline_orientation=False,
            use_seal_recognition=False,
            use_formula_recognition=False,
            use_chart_recognition=False,
            use_region_detection=False,
            use_table_recognition=True,
            text_detection_model_name="PP-OCRv5_mobile_det",
            text_recognition_model_name="PP-OCRv5_mobile_rec",
            # This document is table-heavy; the table-only layout model avoids
            # spending CPU on unrelated document classes and reliably routes
            # detected regions into table recognition.
            layout_detection_model_name="PicoDet_layout_1x_table",
            wired_table_structure_recognition_model_name="SLANet_plus",
            wireless_table_structure_recognition_model_name="SLANet_plus",
            chart_recognition_batch_size=1,
            text_recognition_batch_size=1,
            engine="paddle", enable_mkldnn=False)
    return _ocr, _structure

def blocks(result, page=1):
    data = result.json if hasattr(result, "json") else result
    if isinstance(data, str): data = json.loads(data)
    out=[]
    def walk(x):
        if isinstance(x, dict):
            # PPStructure stores table HTML on table nodes without a `text` field.
            # Preserve it explicitly so downstream chunking can use row/column structure.
            html = x.get("html") or x.get("table_html")
            if html and isinstance(html, str):
                out.append({"block_type":"table","text":"","bbox":x.get("bbox",[]),"confidence":float(x.get("confidence",1.0) or 1.0),"order_index":len(out),"table_html":html})
            text=x.get("text") or x.get("rec_texts")
            score=x.get("confidence", x.get("rec_scores", 0.0))
            box=x.get("bbox") or x.get("rec_boxes") or x.get("dt_polys")
            if isinstance(text, list):
                for i,t in enumerate(text): out.append({"block_type":"text","text":str(t),"bbox":box[i] if isinstance(box,list) and i<len(box) else [],"confidence":float(score[i]) if isinstance(score,list) and i<len(score) else 0.0,"order_index":len(out)})
            elif text: out.append({"block_type":"table" if "table" in x else "text","text":str(text),"bbox":box or [],"confidence":float(score) if isinstance(score,(int,float)) else 0.0,"order_index":len(out)})
            for v in x.values(): walk(v)
        elif isinstance(x,list):
            for v in x: walk(v)
    walk(data)
    return out

@app.get("/health")
def health(): return {"status":"UP"}

@app.post("/v1/ocr")
async def ocr_file(file: UploadFile = File(...), start_page: int = Query(1, ge=1), end_page: int | None = Query(None, ge=1),
                   page: int | None = Query(None, ge=1), crop: str | None = Query(None, description="x0,y0,x1,y1 ratios"),
                   mode: str = Query("structure", pattern="^(structure|plain)$")):
    raw = await file.read(MAX_BYTES + 1)
    if len(raw) > MAX_BYTES: raise HTTPException(413, "file exceeds 100 MB")
    name=(file.filename or "upload").lower(); ext=Path(name).suffix
    if ext not in {".png",".jpg",".jpeg",".bmp",".webp",".pdf",".xlsx"}: raise HTTPException(400,"unsupported file type")
    try:
        ocr, structure = engines(); doc=str(uuid.uuid4()); pages=[]
        if ext == ".pdf":
            import pypdfium2 as pdfium
            pdf=pdfium.PdfDocument(raw)
            first = (page or start_page) - 1
            last = min(len(pdf), page or end_page or len(pdf))
            if first >= last: raise HTTPException(400, "invalid page range")
            for i in range(first, last):
                log.info("processing pdf page %s/%s", i + 1, len(pdf))
                bitmap=pdf[i].render(scale=2); img=normalize_image(np.asarray(bitmap.to_pil().convert("RGB")))
                if crop:
                    try:
                        vals=[float(v.strip()) for v in crop.split(",")]
                        if len(vals)!=4 or any(v<0 or v>1 for v in vals) or vals[2]<=vals[0] or vals[3]<=vals[1]: raise ValueError
                        h,w=img.shape[:2]; x0,y0,x1,y1=vals
                        img=img[int(y0*h):int(y1*h),int(x0*w):int(x1*w)]
                    except ValueError: raise HTTPException(400, "crop must be x0,y0,x1,y1 ratios between 0 and 1")
                if img.size == 0 or img.shape[0] < 32 or img.shape[1] < 32:
                    raise HTTPException(400, "page crop is empty or too small")
                log.info("PPStructure input page=%s shape=%s crop=%s", i + 1, img.shape, crop or "full")
                results=list(ocr.predict(img) if mode == "plain" else structure.predict(img))
                pages.append({"page":i+1,"blocks":blocks(results[0],i+1) if results else []})
        elif ext == ".xlsx":
            import openpyxl
            wb=openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
            for i,ws in enumerate(wb.worksheets,1):
                text="\n".join(" | ".join(str(c.value or "") for c in row) for row in ws.iter_rows())
                pages.append({"page":i,"blocks":[{"block_type":"table","text":text,"bbox":[],"confidence":1.0,"order_index":0}]})
        else:
            img=normalize_image(np.asarray(Image.open(io.BytesIO(raw)).convert("RGB")))
            pages=[{"page":1,"blocks":blocks(next(iter(ocr.predict(img))),1)}]
        return {"document_id":doc,"pages":pages}
    except Exception as e:
        log.exception("OCR failed for %s", name); raise HTTPException(500, str(e))
